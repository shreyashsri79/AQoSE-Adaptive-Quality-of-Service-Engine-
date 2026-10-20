"""The Round 1 one-slider, built on the official template from numbers.json.

  PYTHONPATH=ml .venv/bin/python reports/round1/slide_numbers.py     # numbers from the results
  python reports/round1/build_slide.py [--design N] [--link URL]     # needs python-pptx and qrcode
  python reports/round1/build_slide.py --all                         # every design, into variants/

Slide 1 is the template's cover with the case study number and team from
cover.json. Slide 2 keeps the two logos and nothing else of the template;
everything on it is drawn by one of the designs in designs.py, charts as
native PowerPoint charts so they can be edited. No number is typed into
the designs: each comes from numbers.json, which says where it came from.
"""

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from designs import DESIGNS  # noqa: E402
from slide_kit import cover  # noqa: E402

REPO = HERE.parents[1]
TEMPLATE = REPO / "Teliport S4_Submission Template.pptx"
OUT = HERE / "Round1_OneSlider.pptx"
CHOSEN = 1  # the design the submission uses


def build(design, numbers, info, link, out):
    prs = Presentation(str(TEMPLATE))
    cover(prs.slides[0], info)
    DESIGNS[design][1](prs.slides[1], numbers, link)
    prs.save(out)
    return out


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--link", default="LINK_TBD", help="the Drive link to the detailed PDF")
    parser.add_argument("--design", type=int, default=CHOSEN, choices=sorted(DESIGNS))
    parser.add_argument("--all", action="store_true", help="build every design into variants/")
    args = parser.parse_args(argv)
    numbers = json.loads((HERE / "numbers.json").read_text())
    info_path = HERE / "cover.json"
    info = json.loads(info_path.read_text()) if info_path.exists() else {}
    if args.all:
        (HERE / "variants").mkdir(exist_ok=True)
        jobs = [(d, HERE / "variants" / f"Round1_OneSlider-d{d:02d}.pptx") for d in DESIGNS]
    else:
        jobs = [(args.design, OUT)]
    for design, out in jobs:
        build(design, numbers, info, args.link, out)
        print(out.relative_to(REPO), DESIGNS[design][0])


if __name__ == "__main__":
    main()
